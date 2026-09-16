"""
Document ingestion service.
Supports PDF, PPTX, and DOCX formats.
Extracts raw text, metadata, and structure from uploaded files.
"""

import os
import uuid
import json
import hashlib
from pathlib import Path
from datetime import datetime
from typing import Tuple, Dict, Any

# PDF
import pdfplumber

# PPTX
from pptx import Presentation
from pptx.util import Pt

# DOCX
import docx

from ..models.schemas import DocumentFormat, DocumentMetadata


UPLOAD_DIR = Path(__file__).resolve().parents[2] / "uploads"
DATA_DIR = Path(__file__).resolve().parents[2] / "data"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR.mkdir(parents=True, exist_ok=True)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _count_words(text: str) -> int:
    return len(text.split())


def _detect_format(filename: str) -> DocumentFormat:
    ext = Path(filename).suffix.lower()
    mapping = {
        ".pdf": DocumentFormat.PDF,
        ".pptx": DocumentFormat.PPTX,
        ".docx": DocumentFormat.DOCX,
    }
    if ext not in mapping:
        raise ValueError(f"Unsupported file format: {ext}. Accepted: PDF, PPTX, DOCX.")
    return mapping[ext]


# ─── Extractors ───────────────────────────────────────────────────────────────

def _extract_pdf(filepath: Path) -> Tuple[str, int]:
    """Extract text and page count from a PDF file."""
    pages_text = []
    with pdfplumber.open(str(filepath)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages_text.append(text.strip())
    full_text = "\n\n".join(pages_text)
    return full_text, page_count


def _extract_pptx(filepath: Path) -> Tuple[str, int]:
    """Extract text and slide count from a PPTX file."""
    prs = Presentation(str(filepath))
    slides_text = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_parts = [f"[Slide {slide_num}]"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if len(slide_parts) > 1:
            slides_text.append("\n".join(slide_parts))
    full_text = "\n\n".join(slides_text)
    return full_text, len(prs.slides)


def _extract_docx(filepath: Path) -> Tuple[str, int]:
    """Extract text and estimated page count from a DOCX file."""
    doc = docx.Document(str(filepath))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs)
    # Estimate pages: roughly 300 words per page
    word_count = _count_words(full_text)
    estimated_pages = max(1, round(word_count / 300))
    return full_text, estimated_pages


# ─── Main ingestion function ───────────────────────────────────────────────────

def ingest_document(filepath: Path, original_filename: str) -> Dict[str, Any]:
    """
    Parse a document and return its text content + metadata.

    Returns a dict with keys:
        doc_id, filename, format, text, page_count, word_count,
        uploaded_at, file_size_kb, filepath
    """
    doc_format = _detect_format(original_filename)

    if doc_format == DocumentFormat.PDF:
        text, page_count = _extract_pdf(filepath)
    elif doc_format == DocumentFormat.PPTX:
        text, page_count = _extract_pptx(filepath)
    elif doc_format == DocumentFormat.DOCX:
        text, page_count = _extract_docx(filepath)
    else:
        raise ValueError(f"Unsupported format: {doc_format}")

    if not text.strip():
        raise ValueError(
            "No text could be extracted from the document. "
            "The file may be image-based or password-protected."
        )

    word_count = _count_words(text)
    file_size_kb = round(filepath.stat().st_size / 1024, 2)

    doc_id = str(uuid.uuid4())

    metadata = {
        "doc_id": doc_id,
        "filename": original_filename,
        "format": doc_format.value,
        "page_count": page_count,
        "word_count": word_count,
        "uploaded_at": datetime.utcnow().isoformat(),
        "file_size_kb": file_size_kb,
        "filepath": str(filepath),
    }

    # Persist text + metadata to disk for later retrieval
    doc_dir = DATA_DIR / doc_id
    doc_dir.mkdir(parents=True, exist_ok=True)

    (doc_dir / "text.txt").write_text(text, encoding="utf-8")
    (doc_dir / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    return {**metadata, "text": text}


def load_document(doc_id: str) -> Dict[str, Any]:
    """Load a previously ingested document from disk."""
    doc_dir = DATA_DIR / doc_id
    if not doc_dir.exists():
        raise FileNotFoundError(f"Document '{doc_id}' not found.")

    metadata = json.loads((doc_dir / "metadata.json").read_text(encoding="utf-8"))
    text = (doc_dir / "text.txt").read_text(encoding="utf-8")
    return {**metadata, "text": text}


def list_documents() -> list:
    """List all ingested documents."""
    docs = []
    if DATA_DIR.exists():
        for doc_dir in sorted(DATA_DIR.iterdir()):
            meta_file = doc_dir / "metadata.json"
            if meta_file.exists():
                docs.append(json.loads(meta_file.read_text(encoding="utf-8")))
    return docs
